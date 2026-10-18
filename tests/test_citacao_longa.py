"""Testes da detecção de citações longas (NBR 10520)."""

from __future__ import annotations

import pytest
from docx import Document
from docx.document import Document as DocumentoDocx

from formata_abnt.citacao_longa import (
    aplicar_citacoes_longas,
    eh_citacao_longa,
    estimar_linhas,
    parece_citacao_direta,
)
from formata_abnt.normas import PADRAO
from tests.conftest import CITACAO_LONGA


def test_estimar_linhas() -> None:
    """A estimativa usa ~65 caracteres por linha."""
    assert estimar_linhas("") == 0
    assert estimar_linhas("a" * 65) == 1
    assert estimar_linhas("a" * 66) == 2


def test_parece_citacao_direta() -> None:
    """Reconhece aspas e marcadores como indício de citação."""
    assert parece_citacao_direta('"trecho citado"')
    assert parece_citacao_direta("“trecho citado”")
    assert not parece_citacao_direta("texto comum do autor")


def test_eh_citacao_longa_para_texto_entre_aspas_longo() -> None:
    """Citação entre aspas com mais de 3 linhas é considerada longa."""
    doc = Document()
    paragrafo = doc.add_paragraph(CITACAO_LONGA)

    assert estimar_linhas(CITACAO_LONGA) > 3
    assert eh_citacao_longa(paragrafo, PADRAO)


def test_nao_e_citacao_longa_quando_curta() -> None:
    """Trecho curto entre aspas permanece no corpo."""
    doc = Document()
    paragrafo = doc.add_paragraph('"Trecho curto."')

    assert not eh_citacao_longa(paragrafo, PADRAO)


def test_nao_e_citacao_longa_sem_aspas() -> None:
    """Parágrafo longo sem aspas não é tratado como citação direta."""
    doc = Document()
    paragrafo = doc.add_paragraph("palavra " * 80)

    assert not eh_citacao_longa(paragrafo, PADRAO)


def test_aplicar_citacoes_formata_bloco(documento_exemplo: DocumentoDocx) -> None:
    """Aplica estilo, recuo de 4 cm e fonte 10 pt à citação longa."""
    total = aplicar_citacoes_longas(documento_exemplo, PADRAO)

    assert total == 1
    citacoes = [
        p
        for p in documento_exemplo.paragraphs
        if p.style.name == PADRAO.nome_estilo_citacao
    ]
    assert len(citacoes) == 1
    citacao = citacoes[0]
    assert citacao.paragraph_format.left_indent.cm == pytest.approx(4.0, abs=0.01)
    assert citacao.paragraph_format.line_spacing == 1.0
    assert citacao.runs[0].font.size.pt == 10
