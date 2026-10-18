"""Fixtures compartilhadas dos testes."""

from __future__ import annotations

from collections.abc import Iterator
from pathlib import Path

import pytest
from docx import Document
from docx.document import Document as DocumentoDocx

CITACAO_LONGA = (
    '"A formatação de trabalhos acadêmicos exige atenção rigorosa a detalhes '
    "como margens, tipografia, espaçamento entre linhas, recuo de parágrafo e "
    "apresentação de citações, de modo que o documento final atenda "
    'integralmente às exigências das normas brasileiras vigentes."'
)


@pytest.fixture
def documento_exemplo() -> DocumentoDocx:
    """Documento .docx com título, subtítulo, corpo e citação longa."""
    doc = Document()
    doc.add_heading("1. INTRODUÇÃO", level=1)
    doc.add_paragraph("Parágrafo de corpo com  espaçamento\u00a0irregular.")
    doc.add_heading("1.1 Objetivos", level=2)
    doc.add_paragraph("Segundo parágrafo de corpo do trabalho acadêmico.")
    doc.add_paragraph(CITACAO_LONGA)
    return doc


@pytest.fixture
def arquivo_docx(tmp_path: Path, documento_exemplo: DocumentoDocx) -> Path:
    """Documento .docx gravado em disco, simulando um arquivo real."""
    destino = tmp_path / "entrada.docx"
    documento_exemplo.save(str(destino))
    return destino


@pytest.fixture
def apagar_estilos() -> Iterator[None]:
    """Reserva para testes que precisam de estado limpo de estilos."""
    yield
