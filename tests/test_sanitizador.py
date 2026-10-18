"""Testes da sanitização de texto."""

from __future__ import annotations

from docx import Document
from docx.document import Document as DocumentoDocx
from docx.enum.text import WD_BREAK

from formata_abnt.sanitizador import (
    limpar_runs,
    limpar_texto,
    remover_paragrafos_vazios_repetidos,
    remover_quebras_repetidas,
    sanitizar,
)


def test_limpar_texto_remove_invisiveis_e_espacos() -> None:
    """Remove espaços inseparáveis, hífen suave e colapsa espaços."""
    entrada = "palavra\u00a0com   espacos\u00admuito estranhos"
    assert limpar_texto(entrada) == "palavra com espacosmuito estranhos"


def test_limpar_runs_conta_alterados() -> None:
    """Conta quantos runs foram efetivamente limpos."""
    doc = Document()
    doc.add_paragraph("texto\u00a0normal")
    doc.add_paragraph("texto normal")

    assert limpar_runs(doc) == 1


def test_remove_paragrafos_vazios_repetidos() -> None:
    """Mantém apenas um parágrafo vazio entre blocos de texto."""
    doc = Document()
    doc.add_paragraph("Primeiro")
    doc.add_paragraph("")
    doc.add_paragraph("")
    doc.add_paragraph("Segundo")

    removidos = remover_paragrafos_vazios_repetidos(doc)

    assert removidos == 1
    assert len(doc.paragraphs) == 3


def test_remove_quebras_de_linha_repetidas() -> None:
    """Quebras de linha duplicadas são reduzidas a uma."""
    doc = Document()
    paragrafo = doc.add_paragraph()
    run = paragrafo.add_run("linha")
    run.add_break()
    run.add_break()
    run.add_text("seguinte")

    assert remover_quebras_repetidas(doc) == 1


def test_preserva_quebra_de_pagina() -> None:
    """Quebras de página não são removidas."""
    doc = Document()
    paragrafo = doc.add_paragraph()
    run = paragrafo.add_run("fim")
    run.add_break(WD_BREAK.PAGE)
    run.add_break(WD_BREAK.PAGE)

    assert remover_quebras_repetidas(doc) == 0


def test_sanitizar_agrega_resultados() -> None:
    """A função principal soma runs, parágrafos e quebras limpos."""
    doc: DocumentoDocx = Document()
    doc.add_paragraph("texto\u00a0sujo")
    doc.add_paragraph("")
    doc.add_paragraph("")

    resultado = sanitizar(doc)

    assert resultado.runs_limpos == 1
    assert resultado.paragrafos_removidos == 1
