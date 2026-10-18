"""Testes do modelo POP (cabeçalho e rodapé empresarial)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from formata_abnt.docx_io import abrir, salvar
from formata_abnt.pop import ConfiguracaoPOP, aplicar_modelo_pop, desativar_pop, texto_resumo


def _contem_campo(paragrafo, instrucao: str) -> bool:
    """Verifica se o parágrafo contém um campo Word com a instrução dada."""
    for run in paragrafo.runs:
        for elemento in run._element:
            if elemento.tag == qn("w:instrText") and instrucao in (elemento.text or ""):
                return True
    return False


def test_configuracao_pop_padrao_desativada() -> None:
    """Por padrão o modelo POP está desativado."""
    cfg = ConfiguracaoPOP()

    assert cfg.ativo is False
    assert cfg.classificacao == "USO INTERNO"


def test_aplicar_modelo_pop_desativado_nao_altera() -> None:
    """Quando ativo=False, o documento permanece sem cabeçalho/rodapé."""
    documento = Document()
    secao = documento.sections[0]

    resultado = aplicar_modelo_pop(documento, ConfiguracaoPOP())

    assert resultado.cabecalho_criado is False
    assert len(secao.header.paragraphs) == 1
    assert secao.header.paragraphs[0].text == ""


def test_aplicar_modelo_pop_cria_cabecalho_e_rodape() -> None:
    """O modelo POP cria cabeçalho e rodapé com os campos configurados."""
    documento = Document()
    cfg = ConfiguracaoPOP(
        ativo=True,
        nome_empresa="Minha Empresa",
        titulo_pop="Limpeza de Equipamentos",
        codigo="POP-001",
        revisao="02",
        area_responsavel="Qualidade",
        data_emissao="10/01/2026",
        classificacao="CONFIDENCIAL",
    )

    aplicar_modelo_pop(documento, cfg)

    header_text = documento.sections[0].header.tables[0].rows[0].cells[1].text
    footer_text = documento.sections[0].footer.tables[0].rows[0].cells[0].text

    assert "PROCEDIMENTO OPERACIONAL PADRÃO" in header_text
    assert "Limpeza de Equipamentos" in header_text
    assert "POP-001" in header_text
    assert "CONFIDENCIAL" in footer_text

    # Nome da empresa fica na célula esquerda, junto ao logo.
    empresa_text = documento.sections[0].header.tables[0].rows[0].cells[0].text
    assert "Minha Empresa" in empresa_text


def test_rodape_tem_campos_de_pagina() -> None:
    """O rodapé contém os campos PAGE e NUMPAGES."""
    documento = Document()
    cfg = ConfiguracaoPOP(ativo=True)

    aplicar_modelo_pop(documento, cfg)

    celula_pagina = documento.sections[0].footer.tables[0].rows[0].cells[2]
    assert _contem_campo(celula_pagina.paragraphs[0], "PAGE")
    assert _contem_campo(celula_pagina.paragraphs[0], "NUMPAGES")


def test_desativar_pop_remove_conteudo() -> None:
    """desativar_pop limpa cabeçalho e rodapé."""
    documento = Document()
    cfg = ConfiguracaoPOP(ativo=True, nome_empresa="X")

    aplicar_modelo_pop(documento, cfg)
    desativar_pop(documento)

    assert documento.sections[0].header.tables == []
    assert documento.sections[0].footer.tables == []


def test_texto_resumo_quando_ativo() -> None:
    """O resumo lista os principais campos do modelo."""
    cfg = ConfiguracaoPOP(ativo=True, nome_empresa="X", titulo_pop="Y", codigo="Z")

    resumo = texto_resumo(cfg)

    assert "Cabeçalho POP" in resumo
    assert "X" in resumo
    assert "Y" in resumo
    assert "Z" in resumo
    assert "Página: X de Y" in resumo


def test_texto_resumo_quando_desativado() -> None:
    """O resumo avisa quando o modelo está desativado."""
    assert texto_resumo(ConfiguracaoPOP()) == "Modelo POP desativado."


def test_modelo_pop_persiste_apos_salvar(tmp_path: Path) -> None:
    """O cabeçalho e rodapé POP sobrevivem à gravação do .docx."""
    documento = Document()
    cfg = ConfiguracaoPOP(ativo=True, nome_empresa="Empresa Teste", titulo_pop="POP X")
    aplicar_modelo_pop(documento, cfg)

    destino = tmp_path / "pop.docx"
    salvar(documento, destino)

    reaberto = abrir(destino)
    empresa_text = reaberto.sections[0].header.tables[0].rows[0].cells[0].text
    pop_text = reaberto.sections[0].header.tables[0].rows[0].cells[1].text
    assert "Empresa Teste" in empresa_text
    assert "POP X" in pop_text
