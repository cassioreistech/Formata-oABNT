"""Testes da hierarquia de títulos (NBR 6024)."""

from __future__ import annotations

from docx import Document
from docx.document import Document as DocumentoDocx
from docx.enum.text import WD_ALIGN_PARAGRAPH

from formata_abnt.normas import PADRAO
from formata_abnt.titulos import (
    aplicar_titulos,
    nivel_do_titulo,
    nivel_por_estilo,
    nivel_por_numeracao,
)


def test_nivel_por_numeracao() -> None:
    """Infere o nível pela numeração progressiva do texto."""
    assert nivel_por_numeracao("1. INTRODUÇÃO") == 1
    assert nivel_por_numeracao("1.1 Objetivos") == 2
    assert nivel_por_numeracao("2.3.1 Método") == 3
    assert nivel_por_numeracao("Texto comum sem número") is None


def test_nivel_por_estilo() -> None:
    """Reconhece títulos pelos estilos Heading do Word."""
    doc = Document()
    titulo = doc.add_heading("Introdução", level=2)

    assert nivel_por_estilo(titulo) == 2


def test_nivel_do_titulo_prioriza_estilo() -> None:
    """Um parágrafo com estilo de título é reconhecido."""
    doc = Document()
    titulo = doc.add_heading("Metodologia", level=1)

    assert nivel_do_titulo(titulo) == 1


def test_aplicar_titulos_conta_e_formata(documento_exemplo: DocumentoDocx) -> None:
    """Formata os títulos detectados e devolve a quantidade."""
    total = aplicar_titulos(documento_exemplo, PADRAO)

    assert total == 2
    titulo = documento_exemplo.paragraphs[0]
    assert titulo.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert all(run.bold for run in titulo.runs)


def test_titulo_recebe_estilo_heading(documento_exemplo: DocumentoDocx) -> None:
    """Títulos ganham estilo Heading, útil para o sumário do Word."""
    aplicar_titulos(documento_exemplo, PADRAO)

    assert documento_exemplo.paragraphs[0].style.name == "Heading 1"
    assert documento_exemplo.paragraphs[2].style.name == "Heading 2"
