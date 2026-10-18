"""Testes da aplicação de fonte (NBR 14724)."""

from __future__ import annotations

from docx.document import Document as DocumentoDocx

from formata_abnt.abnt.fontes import aplicar_fonte_corpo, definir_fonte_padrao
from formata_abnt.abnt.normas import NormasABNT, PADRAO
from formata_abnt.abnt.titulos import nivel_do_titulo


def test_fonte_corpo_times_new_roman_12(documento_exemplo: DocumentoDocx) -> None:
    """Runs de corpo recebem Times New Roman tamanho 12."""
    aplicar_fonte_corpo(documento_exemplo, PADRAO)

    corpo = [
        p
        for p in documento_exemplo.paragraphs
        if p.text and nivel_do_titulo(p) is None
    ]
    assert corpo
    for paragrafo in corpo:
        for run in paragrafo.runs:
            assert run.font.name == "Times New Roman"
            assert run.font.size.pt == 12


def test_fonte_corpo_nao_afeta_titulos(documento_exemplo: DocumentoDocx) -> None:
    """A aplicação no corpo não deve definir fonte direta nos títulos."""
    aplicar_fonte_corpo(documento_exemplo, PADRAO)

    titulos = [p for p in documento_exemplo.paragraphs if nivel_do_titulo(p)]
    assert titulos
    for titulo in titulos:
        for run in titulo.runs:
            assert run.font.name is None
            assert run.font.size is None


def test_fonte_customizada(documento_exemplo: DocumentoDocx) -> None:
    """Permite trocar a fonte do corpo via configuração."""
    normas = NormasABNT(fonte_corpo="Arial")
    aplicar_fonte_corpo(documento_exemplo, normas)

    run = documento_exemplo.paragraphs[1].runs[0]
    assert run.font.name == "Arial"


def test_fonte_padrao_do_estilo_normal(documento_exemplo: DocumentoDocx) -> None:
    """O estilo Normal passa a usar a fonte ABNT."""
    definir_fonte_padrao(documento_exemplo, PADRAO)

    estilo = documento_exemplo.styles["Normal"]
    assert estilo.font.name == "Times New Roman"
    assert estilo.font.size.pt == 12
