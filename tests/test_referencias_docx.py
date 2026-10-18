"""Testes da seção REFERÊNCIAS no .docx (NBR 6023)."""

from __future__ import annotations

from docx import Document
from docx.document import Document as DocumentoDocx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm
from pathlib import Path

from abnt_references import ABNTReferenceFormatter, Referencia
from formata_abnt.docx_io import abrir, salvar
from formata_abnt.normas import PADRAO
from formata_abnt.referencias import (
    NOME_SECAO,
    aplicar_negrito_no_intervalo,
    encontrar_secao_referencias,
    escrever_referencia,
    formatar_secao_referencias,
    garantir_secao_referencias,
    inserir_referencia,
    ordenar_referencias,
    paragrafos_da_secao,
    realcar_referencia,
    realcar_secao,
)

FORMATADOR = ABNTReferenceFormatter()


def _livro(autor: str, titulo: str) -> Referencia:
    """Atalho para criar uma referência de livro nos testes."""
    return FORMATADOR.format_book(
        author=autor, title=titulo, city="São Paulo", publisher="Editora", year=2020
    )


def test_garantir_secao_cria_titulo() -> None:
    """A seção REFERÊNCIAS é criada no fim do documento."""
    documento = Document()
    documento.add_paragraph("Texto qualquer.")

    titulo = garantir_secao_referencias(documento)

    assert titulo.text == NOME_SECAO
    assert encontrar_secao_referencias(documento) is not None


def test_garantir_secao_nao_duplica() -> None:
    """Chamar duas vezes não cria duas seções."""
    documento = Document()
    garantir_secao_referencias(documento)
    garantir_secao_referencias(documento)

    titulos = [p for p in documento.paragraphs if p.text.strip() == NOME_SECAO]
    assert len(titulos) == 1


def test_encontrar_secao_ignora_acentos_e_caixa() -> None:
    """A busca aceita 'REFERENCIAS' sem acento."""
    documento = Document()
    documento.add_paragraph("REFERENCES")  # outro idioma não conta
    documento.add_paragraph("referencias")

    assert encontrar_secao_referencias(documento) is not None


def test_escrever_referencia_aplica_negrito() -> None:
    """O título do livro entra em negrito e o parágrafo fica à esquerda."""
    documento = Document()
    paragrafo = documento.add_paragraph()

    escrever_referencia(paragrafo, _livro("Maria Helena Diniz", "Curso de Direito Civil"))

    negritos = [run.text for run in paragrafo.runs if run.bold]
    assert negritos == ["Curso de Direito Civil"]
    assert paragrafo.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert paragrafo.paragraph_format.first_line_indent is None or (
        paragrafo.paragraph_format.first_line_indent.cm == 0
    )
    assert paragrafo.paragraph_format.line_spacing == 1.0
    assert paragrafo.runs[0].font.name == PADRAO.fonte_corpo


def test_inserir_referencia_cria_secao_e_item() -> None:
    """Inserir em documento sem seção cria a seção e adiciona o item."""
    documento = Document()
    documento.add_paragraph("Conteúdo do trabalho.")

    inserir_referencia(documento, _livro("Autor Um", "Primeiro Livro"))

    assert encontrar_secao_referencias(documento) is not None
    assert len(paragrafos_da_secao(documento)) == 1


def test_ordenar_referencias_alfabeticamente() -> None:
    """As referências são reordenadas pelo sobrenome."""
    documento = Document()
    garantir_secao_referencias(documento)
    inserir_referencia(documento, _livro("Zélia Zebra", "Zeta"), ordenar=False)
    inserir_referencia(documento, _livro("Ana Alves", "Alfa"), ordenar=False)

    total = ordenar_referencias(documento)

    textos = [p.text for p in paragrafos_da_secao(documento)]
    assert total == 2
    assert textos[0].startswith("ALVES")
    assert textos[1].startswith("ZEBRA")


def test_ordenar_preserva_negrito() -> None:
    """A reordenação preserva a formatação em negrito dos itens."""
    documento = Document()
    garantir_secao_referencias(documento)
    inserir_referencia(documento, _livro("Zélia Zebra", "Zeta"), ordenar=False)
    inserir_referencia(documento, _livro("Ana Alves", "Alfa"), ordenar=False)

    ordenar_referencias(documento)

    primeiro = paragrafos_da_secao(documento)[0]
    assert [run.text for run in primeiro.runs if run.bold] == ["Alfa"]


def test_formatar_secao_sem_secao_devolve_zero() -> None:
    """Sem seção REFERÊNCIAS nada é alterado."""
    documento = Document()
    documento.add_paragraph("Apenas texto.")

    assert formatar_secao_referencias(documento) == 0


def test_formatar_secao_normaliza_paragrafos() -> None:
    """Parágrafos de referência voltam a ser justificados à esquerda e ordenados."""
    documento = Document()
    documento.add_paragraph(NOME_SECAO)
    zebra = documento.add_paragraph("ZEBRA, Zélia. Zeta. São Paulo: Editora, 2020.")
    zebra.paragraph_format.first_line_indent = Cm(PADRAO.recuo_primeira_linha_cm)
    zebra.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    documento.add_paragraph("ALVES, Ana. Alfa. São Paulo: Editora, 2021.")

    total = formatar_secao_referencias(documento)

    assert total == 2
    itens = paragrafos_da_secao(documento)
    assert itens[0].text.startswith("ALVES")
    assert itens[0].paragraph_format.alignment == WD_ALIGN_PARAGRAPH.LEFT


def test_salvar_e_reabrir_preserva_referencias(tmp_path: Path) -> None:
    """A seção e a ordenação persistem no arquivo gravado."""
    documento: DocumentoDocx = Document()
    inserir_referencia(documento, _livro("Zélia Zebra", "Zeta"))
    inserir_referencia(documento, _livro("Ana Alves", "Alfa"))

    destino = tmp_path / "com_referencias.docx"
    salvar(documento, destino)

    reaberto = abrir(destino)
    itens = paragrafos_da_secao(reaberto)
    assert len(itens) == 2
    assert itens[0].text.startswith("ALVES")
    assert any(run.bold for run in itens[0].runs)


# ── Reanálise automática do negrito ─────────────────────────────────────────


def test_aplicar_negrito_divide_run_ao_meio() -> None:
    """Um run único é fatiado nas fronteiras do intervalo, sem perder texto."""
    documento = Document()
    paragrafo = documento.add_paragraph("abcdefghij")

    assert aplicar_negrito_no_intervalo(paragrafo, 3, 7)

    assert paragrafo.text == "abcdefghij"
    assert [run.text for run in paragrafo.runs] == ["abc", "defg", "hij"]
    assert [run.text for run in paragrafo.runs if run.bold] == ["defg"]


def test_realcar_referencia_preserva_texto_e_destaca_livro() -> None:
    """O texto do livro fica idêntico; só o título entra em negrito."""
    original = (
        "DINIZ, Maria Helena. Curso de Direito Civil. 38 ed. São Paulo: Saraiva, 2022."
    )
    documento = Document()
    paragrafo = documento.add_paragraph(original)

    assert realcar_referencia(paragrafo) is True

    assert paragrafo.text == original
    assert [run.text for run in paragrafo.runs if run.bold] == ["Curso de Direito Civil"]


def test_realcar_referencia_destaca_periodico() -> None:
    """Em artigos, o negrito vai para o título do periódico."""
    original = (
        "SANTOS, Carlos. Impactos da IA no Direito. Revista Brasileira de Tecnologia, "
        "Brasília, v. 12, n. 2, p. 45-60, 2024."
    )
    documento = Document()
    paragrafo = documento.add_paragraph(original)

    realcar_referencia(paragrafo)

    assert paragrafo.text == original
    assert [run.text for run in paragrafo.runs if run.bold] == [
        "Revista Brasileira de Tecnologia"
    ]


def test_realcar_referencia_remove_negrito_fora_do_lugar() -> None:
    """Negrito aplicado no lugar errado é corrigido."""
    original = "ALVES, Ana. Alfa. São Paulo: Editora, 2021."
    documento = Document()
    paragrafo = documento.add_paragraph()
    paragrafo.add_run("ALVES, Ana. ").bold = True
    paragrafo.add_run("Alfa. São Paulo: Editora, 2021.")

    realcar_referencia(paragrafo)

    assert paragrafo.text == original
    assert [run.text for run in paragrafo.runs if run.bold] == ["Alfa"]


def test_realcar_referencia_generica_nao_altera() -> None:
    """Referência genérica não tem destaque aplicado."""
    documento = Document()
    paragrafo = documento.add_paragraph("Anotação solta sem padrão ABNT")

    assert realcar_referencia(paragrafo) is False
    assert paragrafo.text == "Anotação solta sem padrão ABNT"


def test_realcar_secao_conta_apenas_alteradas() -> None:
    """A primeira passada altera; a segunda não encontra o que mudar."""
    documento = Document()
    garantir_secao_referencias(documento)
    documento.add_paragraph("ALVES, Ana. Alfa. São Paulo: Editora, 2021.")
    documento.add_paragraph("ZEBRA, Zélia. Zeta. São Paulo: Editora, 2020.")

    assert realcar_secao(documento) == 2
    assert realcar_secao(documento) == 0


def test_formatar_secao_realca_e_ordena() -> None:
    """A rotina completa realça o negrito e reordena alfabeticamente."""
    documento = Document()
    garantir_secao_referencias(documento)
    documento.add_paragraph("ZEBRA, Zélia. Zeta. São Paulo: Editora, 2020.")
    documento.add_paragraph("ALVES, Ana. Alfa. São Paulo: Editora, 2021.")

    total = formatar_secao_referencias(documento)

    assert total == 2
    itens = paragrafos_da_secao(documento)
    assert itens[0].text.startswith("ALVES")
    assert [run.text for run in itens[0].runs if run.bold] == ["Alfa"]
    assert [run.text for run in itens[1].runs if run.bold] == ["Zeta"]


def test_formatar_secao_pode_desativar_realce() -> None:
    """Com realcar=False, só a formatação e a ordem mudam."""
    documento = Document()
    garantir_secao_referencias(documento)
    documento.add_paragraph("ZEBRA, Zélia. Zeta. São Paulo: Editora, 2020.")
    documento.add_paragraph("ALVES, Ana. Alfa. São Paulo: Editora, 2021.")

    formatar_secao_referencias(documento, realcar=False)

    itens = paragrafos_da_secao(documento)
    assert itens[0].text.startswith("ALVES")
    assert not any(run.bold for run in itens[0].runs)


def test_realce_persiste_apos_salvar_e_reabrir(tmp_path: Path) -> None:
    """O negrito realçado continua no arquivo gravado."""
    documento: DocumentoDocx = Document()
    garantir_secao_referencias(documento)
    documento.add_paragraph(
        "DINIZ, Maria Helena. Curso de Direito Civil. 38 ed. São Paulo: Saraiva, 2022."
    )

    destino = tmp_path / "realcado.docx"
    formatar_secao_referencias(documento)
    salvar(documento, destino)

    reaberto = abrir(destino)
    item = paragrafos_da_secao(reaberto)[0]
    assert item.text == (
        "DINIZ, Maria Helena. Curso de Direito Civil. 38 ed. São Paulo: Saraiva, 2022."
    )
    assert [run.text for run in item.runs if run.bold] == ["Curso de Direito Civil"]
