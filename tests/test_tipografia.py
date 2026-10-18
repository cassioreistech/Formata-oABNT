"""Testes da tipografia do corpo do texto."""

from __future__ import annotations

import pytest
from docx import Document
from docx.document import Document as DocumentoDocx
from docx.enum.text import WD_ALIGN_PARAGRAPH

from formata_abnt.citacao_longa import aplicar_citacoes_longas
from formata_abnt.normas import PADRAO
from formata_abnt.tipografia import aplicar_tipografia_corpo, definir_estilo_normal


def test_aplica_fonte_espacamento_recuo_justificado(
    documento_exemplo: DocumentoDocx,
) -> None:
    """Corpo fica Arial 12, 1,5, justificado e recuo de 1,25 cm."""
    aplicar_tipografia_corpo(documento_exemplo, PADRAO)

    corpo = documento_exemplo.paragraphs[1]
    formato = corpo.paragraph_format
    assert formato.line_spacing == 1.5
    assert formato.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert formato.first_line_indent.cm == pytest.approx(1.25, abs=0.01)
    assert corpo.runs[0].font.name == "Arial"
    assert corpo.runs[0].font.size.pt == 12


def test_preserva_negrito_italico_sublinhado() -> None:
    """A normalização de fonte não deve apagar negrito/itálico/sublinhado."""
    doc = Document()
    paragrafo = doc.add_paragraph()
    run = paragrafo.add_run("texto marcado")
    run.bold = True
    run.italic = True
    run.underline = True

    aplicar_tipografia_corpo(doc, PADRAO)

    assert run.bold is True
    assert run.italic is True
    assert run.underline is True
    assert run.font.name == "Arial"


def test_nao_altera_titulos(documento_exemplo: DocumentoDocx) -> None:
    """Títulos não recebem o recuo de primeira linha do corpo."""
    aplicar_tipografia_corpo(documento_exemplo, PADRAO)

    titulo = documento_exemplo.paragraphs[0]
    assert titulo.paragraph_format.first_line_indent is None


def test_ignora_paragrafos_de_citacao(documento_exemplo: DocumentoDocx) -> None:
    """Parágrafos marcados como citação longa não são tratados como corpo."""
    aplicar_citacoes_longas(documento_exemplo, PADRAO)

    total = aplicar_tipografia_corpo(documento_exemplo, PADRAO)

    assert total == 2  # apenas os dois parágrafos de corpo


def test_definir_estilo_normal(documento_exemplo: DocumentoDocx) -> None:
    """O estilo Normal herda a fonte ABNT."""
    definir_estilo_normal(documento_exemplo, PADRAO)

    estilo = documento_exemplo.styles["Normal"]
    assert estilo.font.name == "Arial"
    assert estilo.font.size.pt == 12
