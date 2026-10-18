"""Testes de ponta a ponta do motor ABNT."""

from __future__ import annotations

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path

from abnt_engine import MotorABNT, formatar_para_abnt
from formata_abnt.docx_io import abrir
from formata_abnt.normas import PADRAO
from tests.conftest import CITACAO_LONGA


@pytest.fixture
def arquivo_sujo(tmp_path: Path) -> Path:
    """Arquivo .docx com títulos, citação longa e sujeira de cópia."""
    doc = Document()
    doc.add_paragraph("1. INTRODUÇÃO")
    doc.add_paragraph("Texto com  espaços\u00a0irregulares e artefatos.")
    doc.add_paragraph("")
    doc.add_paragraph("")
    doc.add_paragraph(CITACAO_LONGA)
    doc.add_paragraph("Conclusão do trabalho acadêmico.")
    destino = tmp_path / "sujo.docx"
    doc.save(str(destino))
    return destino


def test_processar_gera_arquivo(arquivo_sujo: Path, tmp_path: Path) -> None:
    """O pipeline gera o arquivo de saída."""
    destino = tmp_path / "limpo_abnt.docx"

    resultado = formatar_para_abnt(arquivo_sujo, destino)

    assert resultado.arquivo_saida == destino
    assert destino.exists()


def test_processar_aplica_regras(arquivo_sujo: Path, tmp_path: Path) -> None:
    """Títulos, citações e sanitização são contabilizados."""
    resultado = MotorABNT().processar(arquivo_sujo, tmp_path / "saida.docx")

    assert resultado.titulos_aplicados >= 1
    assert resultado.citacoes_aplicadas == 1
    assert resultado.paragrafos_removidos >= 1
    assert resultado.runs_sanitizados >= 1


def test_resultado_persiste_apos_reabrir(arquivo_sujo: Path, tmp_path: Path) -> None:
    """As regras permanecem após salvar e reabrir o documento."""
    destino = tmp_path / "saida.docx"
    formatar_para_abnt(arquivo_sujo, destino)

    reaberto = abrir(destino)
    secao = reaberto.sections[0]
    assert secao.left_margin.cm == pytest.approx(3.0, abs=0.01)
    assert secao.top_margin.cm == pytest.approx(3.0, abs=0.01)

    corpo = reaberto.paragraphs[1]
    assert corpo.paragraph_format.line_spacing == 1.5
    assert corpo.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert corpo.runs[0].font.name == PADRAO.fonte_corpo


def test_notifica_progresso(arquivo_sujo: Path, tmp_path: Path) -> None:
    """O callback de progresso recebe as etapas do pipeline."""
    mensagens: list[str] = []

    MotorABNT().processar(
        arquivo_sujo, tmp_path / "saida.docx", progresso=mensagens.append
    )

    assert mensagens[0] == "Lendo documento..."
    assert mensagens[-1] == "Concluído!"
    assert any("margens" in m.lower() for m in mensagens)


def test_nome_de_saida_automatico(arquivo_sujo: Path) -> None:
    """Sem saída informada, o arquivo recebe o sufixo _abnt."""
    resultado = MotorABNT().processar(arquivo_sujo)

    assert resultado.arquivo_saida.name == "sujo_abnt.docx"


def test_pipeline_chama_formatacao_padrao(tmp_path: Path) -> None:
    """O pipeline padrão formata o documento sem POP."""
    documento: DocumentoDocx = Document()
    documento.add_paragraph("1. INTRODUCAO")
    documento.add_paragraph("Texto do capitulo.")

    entrada = tmp_path / "doc.docx"
    saida = tmp_path / "doc_abnt.docx"
    documento.save(str(entrada))

    resultado = formatar_para_abnt(entrada, saida)

    assert resultado.arquivo_saida == saida
    assert resultado.paragrafos_formatados > 0
    assert resultado.pop_aplicado is False


def test_pipeline_aplica_pop_quando_preferencia_ativa(tmp_path: Path) -> None:
    """O pipeline aplica o modelo POP quando a preferência está ativa."""
    from formata_abnt.config import ConfiguracaoPOP, Preferencias

    documento: DocumentoDocx = Document()
    documento.add_paragraph("1. INTRODUCAO")

    entrada = tmp_path / "doc.docx"
    saida = tmp_path / "doc_abnt.docx"
    documento.save(str(entrada))

    preferencias = Preferencias(
        pop=ConfiguracaoPOP(ativo=True, nome_empresa="Empresa", titulo_pop="POP X")
    )
    resultado = formatar_para_abnt(entrada, saida, preferencias=preferencias)

    assert resultado.pop_aplicado is True

    reaberto = abrir(saida)
    pop_text = reaberto.sections[0].header.tables[0].rows[0].cells[1].text
    empresa_text = reaberto.sections[0].header.tables[0].rows[0].cells[0].text
    assert "Empresa" in empresa_text
    assert "POP X" in pop_text


def test_pipeline_nao_aplica_pop_quando_preferencia_inativa(tmp_path: Path) -> None:
    """O pipeline não aplica POP quando a preferência está desativada."""
    from formata_abnt.config import ConfiguracaoPOP, Preferencias

    documento: DocumentoDocx = Document()
    entrada = tmp_path / "doc.docx"
    saida = tmp_path / "doc_abnt.docx"
    documento.save(str(entrada))

    preferencias = Preferencias(pop=ConfiguracaoPOP(ativo=False))
    resultado = formatar_para_abnt(entrada, saida, preferencias=preferencias)

    assert resultado.pop_aplicado is False
    assert abrir(saida).sections[0].header.tables == []


def test_engine_ordena_referencias(tmp_path: Path) -> None:
    """O pipeline ordena alfabeticamente as referências da seção REFERÊNCIAS."""
    from abnt_references import ABNTReferenceFormatter
    from formata_abnt.referencias import inserir_referencia, paragrafos_da_secao

    formatador = ABNTReferenceFormatter()
    documento = Document()
    documento.add_paragraph("1. INTRODUÇÃO")
    documento.add_paragraph("Texto do trabalho.")
    inserir_referencia(
        documento,
        formatador.format_book(author="Zélia Zebra", title="Zeta", city="S", publisher="P", year=2020),
        ordenar=False,
    )
    inserir_referencia(
        documento,
        formatador.format_book(author="Ana Alves", title="Alfa", city="S", publisher="P", year=2021),
        ordenar=False,
    )

    origem = tmp_path / "refs.docx"
    documento.save(str(origem))

    resultado = MotorABNT().processar(origem, tmp_path / "refs_abnt.docx")

    assert resultado.referencias_ordenadas == 2
    reaberto = abrir(resultado.arquivo_saida)
    itens = paragrafos_da_secao(reaberto)
    assert itens[0].text.startswith("ALVES")
    assert itens[1].text.startswith("ZEBRA")
    # O pipeline também reaplica o negrito ABNT nos títulos.
    assert [run.text for run in itens[0].runs if run.bold] == ["Alfa"]
    assert [run.text for run in itens[1].runs if run.bold] == ["Zeta"]
