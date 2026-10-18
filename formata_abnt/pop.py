"""Modelo POP (Procedimento Operacional Padrão) para cabeçalho e rodapé.

Ativado apenas quando o usuário marca a opção correspondente, para não misturar
com a formatação ABNT comum de textos acadêmicos.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from docx.document import Document as DocumentoDocx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.table import Table

PADRAO_CLASSIFICACAO = "USO INTERNO"


@dataclass
class ConfiguracaoPOP:
    """Configuração do cabeçalho/rodapé no modelo POP/empresarial."""

    ativo: bool = False
    nome_empresa: str = ""
    logo: str = ""
    titulo_pop: str = ""
    codigo: str = ""
    revisao: str = ""
    area_responsavel: str = ""
    elaborado_por: str = ""
    aprovado_por: str = ""
    data_emissao: str = ""
    classificacao: str = PADRAO_CLASSIFICACAO


@dataclass
class SecaoPOP:
    """Resumo da aplicação do modelo POP em um documento."""

    cabecalho_criado: bool = False
    rodape_criado: bool = False


# ── Utilitários XML ─────────────────────────────────────────────────────────


def _adicionar_campo(paragrafo, instrucao: str, exibicao: str = "1") -> None:
    """Insere um campo de Word (PAGE, NUMPAGES etc.) em um parágrafo."""
    # begin
    run = paragrafo.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    run._r.append(fld_char)

    # instrução
    run = paragrafo.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instrucao} "
    run._r.append(instr)

    # separate
    run = paragrafo.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run._r.append(sep)

    # texto exibido (Word atualiza ao abrir)
    paragrafo.add_run(exibicao)

    # end
    run = paragrafo.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def _formatar_run(run, tamanho: int = 8, negrito: bool = False) -> None:
    """Define a fonte padrão do modelo POP."""
    run.font.name = "Arial"
    run.font.size = Pt(tamanho)
    run.bold = negrito


def _formatar_paragrafo(paragrafo, alinhamento=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    """Aplica alinhamento e remove espaçamento do parágrafo."""
    paragrafo.alignment = alinhamento
    formato = paragrafo.paragraph_format
    formato.space_before = Pt(0)
    formato.space_after = Pt(0)
    formato.line_spacing = 1.0


def _limpar_secao(secao) -> None:
    """Remove o conteúdo existente de um cabeçalho ou rodapé."""
    secao.is_linked_to_previous = False
    for paragrafo in secao.paragraphs:
        paragrafo.clear()
    for tabela in list(secao.tables):
        tabela._element.getparent().remove(tabela._element)


def _borda_inferior(tabela: Table, tamanho: int = 8, cor: str = "000000") -> None:
    """Deixa só a borda inferior da tabela visível."""
    tbl_pr = tabela._tbl.tblPr
    tbl_borders = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{lado}")
        if lado == "bottom":
            border.set(qn("w:val"), "single")
            border.set(qn("w:sz"), str(tamanho))
            border.set(qn("w:color"), cor)
        else:
            border.set(qn("w:val"), "nil")
        tbl_borders.append(border)
    tbl_pr.append(tbl_borders)


def _adicionar_logo_ou_placeholder(cell, config: ConfiguracaoPOP) -> None:
    """Insere o logo da empresa ou um placeholder textual."""
    paragrafo = cell.paragraphs[0]
    _formatar_paragrafo(paragrafo, WD_ALIGN_PARAGRAPH.LEFT)

    caminho_logo = Path(config.logo) if config.logo else None
    if caminho_logo and caminho_logo.is_file():
        run = paragrafo.add_run()
        try:
            run.add_picture(str(caminho_logo), width=Inches(0.8))
            return
        except Exception:  # noqa: BLE001 - fallback em caso de imagem inválida
            pass

    run = paragrafo.add_run("LOGO")
    _formatar_run(run, tamanho=10, negrito=True)

    if config.nome_empresa:
        paragrafo.add_run("\n")
        run = paragrafo.add_run(config.nome_empresa)
        _formatar_run(run, tamanho=8, negrito=True)


def _linha_info(paragrafo, rotulo: str, valor: str) -> None:
    """Adiciona uma linha 'Rótulo: valor' no cabeçalho POP."""
    if not valor:
        return
    run = paragrafo.add_run(f"{rotulo}: ")
    _formatar_run(run, tamanho=8, negrito=True)
    run = paragrafo.add_run(valor)
    _formatar_run(run, tamanho=8)
    paragrafo.add_run("\n")


# ── Aplicação do modelo ─────────────────────────────────────────────────────


def _aplicar_cabecalho(documento: DocumentoDocx, config: ConfiguracaoPOP) -> None:
    """Cria o cabeçalho POP: logo/empresa à esquerda, dados do POP à direita."""
    secao = documento.sections[0]
    cabecalho = secao.header
    _limpar_secao(cabecalho)

    tabela = cabecalho.add_table(rows=1, cols=2, width=Inches(7.0))
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabela.autofit = False
    tabela.allow_autofit = False
    _borda_inferior(tabela)

    colunas = tabela.columns
    colunas[0].width = Inches(2.5)
    colunas[1].width = Inches(4.5)

    celula_esquerda, celula_direita = tabela.rows[0].cells
    _adicionar_logo_ou_placeholder(celula_esquerda, config)

    paragrafo = celula_direita.paragraphs[0]
    _formatar_paragrafo(paragrafo, WD_ALIGN_PARAGRAPH.RIGHT)

    run = paragrafo.add_run("PROCEDIMENTO OPERACIONAL PADRÃO")
    _formatar_run(run, tamanho=9, negrito=True)
    paragrafo.add_run("\n")

    if config.titulo_pop:
        run = paragrafo.add_run(config.titulo_pop)
        _formatar_run(run, tamanho=8, negrito=True)
        paragrafo.add_run("\n")

    _linha_info(paragrafo, "Código", config.codigo)
    _linha_info(paragrafo, "Revisão", config.revisao)
    _linha_info(paragrafo, "Área", config.area_responsavel)

    # Remove quebra extra no final do parágrafo direito.
    if paragrafo.runs and paragrafo.runs[-1].text.endswith("\n"):
        paragrafo.runs[-1].text = paragrafo.runs[-1].text.rstrip("\n")


def _aplicar_rodape(documento: DocumentoDocx, config: ConfiguracaoPOP) -> None:
    """Cria o rodapé POP: classificação, data e numeração de páginas."""
    secao = documento.sections[0]
    rodape = secao.footer
    _limpar_secao(rodape)

    tabela = rodape.add_table(rows=1, cols=3, width=Inches(7.0))
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabela.autofit = False
    tabela.allow_autofit = False
    _borda_inferior(tabela, tamanho=4)

    for coluna in tabela.columns:
        coluna.width = Inches(2.33)

    esquerda, centro, direita = tabela.rows[0].cells

    paragrafo = esquerda.paragraphs[0]
    _formatar_paragrafo(paragrafo, WD_ALIGN_PARAGRAPH.LEFT)
    if config.classificacao:
        run = paragrafo.add_run(config.classificacao)
        _formatar_run(run, tamanho=8, negrito=True)

    paragrafo = centro.paragraphs[0]
    _formatar_paragrafo(paragrafo, WD_ALIGN_PARAGRAPH.CENTER)
    _linha_info(paragrafo, "Emissão", config.data_emissao)
    if paragrafo.runs and paragrafo.runs[-1].text.endswith("\n"):
        paragrafo.runs[-1].text = paragrafo.runs[-1].text.rstrip("\n")

    paragrafo = direita.paragraphs[0]
    _formatar_paragrafo(paragrafo, WD_ALIGN_PARAGRAPH.RIGHT)
    run = paragrafo.add_run("Página ")
    _formatar_run(run, tamanho=8)
    _adicionar_campo(paragrafo, "PAGE")
    run = paragrafo.add_run(" de ")
    _formatar_run(run, tamanho=8)
    _adicionar_campo(paragrafo, "NUMPAGES")


def aplicar_modelo_pop(
    documento: DocumentoDocx, config: ConfiguracaoPOP | None = None
) -> SecaoPOP:
    """Aplica o cabeçalho e rodapé do modelo POP ao documento.

    Se ``config`` for None ou ``ativo=False``, nada é alterado.
    """
    if config is None or not config.ativo:
        return SecaoPOP()

    _aplicar_cabecalho(documento, config)
    _aplicar_rodape(documento, config)
    return SecaoPOP(cabecalho_criado=True, rodape_criado=True)


def desativar_pop(documento: DocumentoDocx) -> None:
    """Remove o conteúdo do cabeçalho e rodapé e os vincula à seção anterior."""
    for secao in documento.sections:
        _limpar_secao(secao.header)
        _limpar_secao(secao.footer)
        secao.header.is_linked_to_previous = True
        secao.footer.is_linked_to_previous = True


def texto_resumo(config: ConfiguracaoPOP) -> str:
    """Devolve um resumo textual do modelo POP para preview."""
    if not config.ativo:
        return "Modelo POP desativado."

    linhas = ["Cabeçalho POP", "─" * 30]
    if config.nome_empresa:
        linhas.append(f"Empresa: {config.nome_empresa}")
    if config.titulo_pop:
        linhas.append(f"Título: {config.titulo_pop}")
    if config.codigo:
        linhas.append(f"Código: {config.codigo}")
    if config.revisao:
        linhas.append(f"Revisão: {config.revisao}")
    if config.area_responsavel:
        linhas.append(f"Área: {config.area_responsavel}")

    linhas.extend(["", "Rodapé POP", "─" * 30])
    linhas.append(f"Classificação: {config.classificacao or '-'}")
    linhas.append(f"Emissão: {config.data_emissao or '-'}")
    linhas.append("Página: X de Y")
    return "\n".join(linhas)
