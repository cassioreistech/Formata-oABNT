"""Seção REFERÊNCIAS em documentos .docx (ABNT NBR 6023).

Responsável por localizar/criar a seção REFERÊNCIAS, inserir referências
com o destaque em negrito real e ordenar automaticamente os itens em
ordem alfabética no final do arquivo.
"""

from __future__ import annotations

import copy
import re
import unicodedata
from collections.abc import Iterable

from docx.document import Document as DocumentoDocx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from abnt_references import ABNTReferenceFormatter, Referencia
from formata_abnt.normas import PADRAO, NormasABNT
from formata_abnt.runs import definir_fonte

NOME_SECAO = "REFERÊNCIAS"

_FORMATADOR = ABNTReferenceFormatter()


def normalizar(texto: str) -> str:
    """Minúsculas, sem acentos e sem espaços excedentes (para comparação)."""
    decomposto = unicodedata.normalize("NFKD", texto)
    sem_acento = "".join(c for c in decomposto if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", sem_acento).strip().casefold()


def eh_titulo_secao(paragrafo: Paragraph) -> bool:
    """Indica se o parágrafo é o título da seção REFERÊNCIAS."""
    return normalizar(paragrafo.text) == normalizar(NOME_SECAO)


def encontrar_secao_referencias(documento: DocumentoDocx) -> Paragraph | None:
    """Devolve o parágrafo do título REFERÊNCIAS, se existir."""
    for paragrafo in documento.paragraphs:
        if eh_titulo_secao(paragrafo):
            return paragrafo
    return None


def garantir_secao_referencias(documento: DocumentoDocx) -> Paragraph:
    """Devolve o título REFERÊNCIAS, criando a seção no fim do documento."""
    existente = encontrar_secao_referencias(documento)
    if existente is not None:
        return existente

    paragrafos = documento.paragraphs
    if paragrafos and paragrafos[-1].text.strip():
        documento.add_page_break()
    return documento.add_heading(NOME_SECAO, level=1)


def definir_formatacao_referencia(paragrafo: Paragraph, normas: NormasABNT) -> None:
    """Aplica o padrão ABNT de referência: alinhado à esquerda, espaço simples."""
    formato = paragrafo.paragraph_format
    formato.alignment = WD_ALIGN_PARAGRAPH.LEFT
    formato.first_line_indent = Cm(0)
    formato.left_indent = Cm(0)
    formato.line_spacing = normas.espacamento_citacao
    formato.space_before = Pt(0)
    formato.space_after = Pt(12)


def escrever_referencia(
    paragrafo: Paragraph, referencia: Referencia, normas: NormasABNT = PADRAO
) -> Paragraph:
    """Escreve os segmentos da referência no parágrafo, com negrito real."""
    definir_formatacao_referencia(paragrafo, normas)

    for run in list(paragrafo.runs):
        run._element.getparent().remove(run._element)

    for segmento in referencia.segmentos:
        if not segmento.texto:
            continue
        run = paragrafo.add_run(segmento.texto)
        definir_fonte(run, normas.fonte_corpo, normas.tamanho_corpo_pt)
        run.bold = segmento.negrito

    return paragrafo


def _limites_dos_runs(paragrafo: Paragraph) -> list[tuple[Run, int, int]]:
    """Devolve ``[(run, início, fim)]`` com os deslocamentos no texto do parágrafo."""
    limites: list[tuple[Run, int, int]] = []
    deslocamento = 0
    for run in paragrafo.runs:
        tamanho = len(run.text)
        limites.append((run, deslocamento, deslocamento + tamanho))
        deslocamento += tamanho
    return limites


def _dividir_run(run: Run, posicao: int) -> Run:
    """Divide o run em ``posicao`` e devolve o novo run com a parte final.

    Clona o elemento XML, portanto a formatação é mantida nas duas metades.
    """
    texto = run.text
    novo_elemento = copy.deepcopy(run._element)
    run._element.addnext(novo_elemento)

    novo_run = Run(novo_elemento, run._parent)
    run.text = texto[:posicao]
    novo_run.text = texto[posicao:]
    return novo_run


def _dividir_no_deslocamento(paragrafo: Paragraph, deslocamento: int) -> None:
    """Garante que ``deslocamento`` seja fronteira entre dois runs."""
    for run, inicio, fim in _limites_dos_runs(paragrafo):
        if inicio < deslocamento < fim:
            _dividir_run(run, deslocamento - inicio)
            return


def aplicar_negrito_no_intervalo(
    paragrafo: Paragraph, inicio: int, fim: int
) -> bool:
    """Marca em negrito o intervalo informado, preservando todo o texto.

    Os runs são divididos nas fronteiras do intervalo; nenhum caractere é
    adicionado, removido ou reordenado.
    """
    if inicio >= fim:
        return False

    _dividir_no_deslocamento(paragrafo, fim)
    _dividir_no_deslocamento(paragrafo, inicio)

    alterou = False
    for run, run_inicio, run_fim in _limites_dos_runs(paragrafo):
        dentro = run_inicio >= inicio and run_fim <= fim
        if dentro and run.bold is not True:
            run.bold = True
            alterou = True
        elif not dentro and run.bold:
            run.bold = False
            alterou = True
    return alterou


def realcar_referencia(paragrafo: Paragraph) -> bool:
    """Reaplica o negrito ABNT em uma referência já escrita, sem alterar o texto.

    Devolve ``True`` se a formatação de destaque mudou.
    """
    intervalo = _FORMATADOR.localizar_destaque(paragrafo.text)
    if intervalo is None:
        return False
    return aplicar_negrito_no_intervalo(paragrafo, intervalo[0], intervalo[1])


def realcar_secao(documento: DocumentoDocx) -> int:
    """Reaplica o negrito em todas as referências e devolve quantas mudaram."""
    alteradas = 0
    for paragrafo in paragrafos_da_secao(documento):
        if realcar_referencia(paragrafo):
            alteradas += 1
    return alteradas


def paragrafos_da_secao(documento: DocumentoDocx) -> list[Paragraph]:
    """Lista os parágrafos de referência (não vazios) após o título."""
    titulo = encontrar_secao_referencias(documento)
    if titulo is None:
        return []

    paragrafos = documento.paragraphs
    alvo = titulo._element
    for indice, paragrafo in enumerate(paragrafos):
        if paragrafo._element is alvo:
            return [p for p in paragrafos[indice + 1 :] if p.text.strip()]
    return []


def ordenar_referencias(documento: DocumentoDocx) -> int:
    """Ordena alfabeticamente as referências da seção e devolve a quantidade."""
    referencias = paragrafos_da_secao(documento)
    if len(referencias) < 2:
        return len(referencias)

    ordenadas = sorted(referencias, key=lambda p: normalizar(p.text))

    corpo = documento.element.body
    sect_pr = corpo.find(qn("w:sectPr"))

    for paragrafo in referencias:
        corpo.remove(paragrafo._element)
    for paragrafo in ordenadas:
        if sect_pr is not None:
            sect_pr.addprevious(paragrafo._element)
        else:
            corpo.append(paragrafo._element)

    return len(ordenadas)


def inserir_referencia(
    documento: DocumentoDocx,
    referencia: Referencia,
    normas: NormasABNT = PADRAO,
    ordenar: bool = True,
) -> Paragraph:
    """Insere a referência ao final da seção REFERÊNCIAS do documento."""
    garantir_secao_referencias(documento)
    paragrafo = escrever_referencia(documento.add_paragraph(), referencia, normas)

    if ordenar:
        ordenar_referencias(documento)
    return paragrafo


def inserir_referencias(
    documento: DocumentoDocx,
    referencias: Iterable[Referencia],
    normas: NormasABNT = PADRAO,
) -> int:
    """Insere várias referências de uma vez e ordena só uma vez no final.

    Usado pela lista de referências em preparo da interface gráfica.
    """
    pendentes = [referencia for referencia in referencias if referencia.segmentos]
    if not pendentes:
        return 0

    garantir_secao_referencias(documento)
    for referencia in pendentes:
        escrever_referencia(documento.add_paragraph(), referencia, normas)

    ordenar_referencias(documento)
    return len(pendentes)


def formatar_secao_referencias(
    documento: DocumentoDocx, normas: NormasABNT = PADRAO, realcar: bool = True
) -> int:
    """Reformata, reaplica o negrito ABNT e ordena as referências existentes.

    Não cria a seção caso ela não exista; nesse caso devolve ``0``.
    """
    referencias = paragrafos_da_secao(documento)
    if not referencias:
        return 0

    for paragrafo in referencias:
        definir_formatacao_referencia(paragrafo, normas)
        if realcar:
            realcar_referencia(paragrafo)
        for run in paragrafo.runs:
            definir_fonte(run, normas.fonte_corpo, normas.tamanho_corpo_pt)

    ordenar_referencias(documento)
    return len(referencias)
